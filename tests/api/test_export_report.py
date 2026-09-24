"""E8 chart export: the training report data and its PNG / PDF / Excel renderers."""

from __future__ import annotations

import pytest
from helpers.runs import completed_fake_run
from PIL import Image

from horos.api.export import export_training_report, list_exports
from horos.api.report import PNG_SIZE, build_training_report
from horos.errors import ProjectError

pytest.importorskip("matplotlib")
pytest.importorskip("openpyxl")


@pytest.fixture(scope="module")
def run(tmp_path_factory):
    return completed_fake_run(tmp_path_factory.mktemp("run"), epochs=3)


def test_report_gathers_the_run(run):
    project, record = run
    report = build_training_report(project, record.run_id)
    assert report.run_id == record.run_id and report.state == "completed"
    assert report.model == "rfdetr-nano" and report.license == "Apache-2.0"  # R3
    assert report.epochs_planned == 3 and report.epochs_completed == 3
    assert report.dataset_images == 3 and report.dataset_splits["train"] == 2
    assert report.class_instances == {"forklift": 2, "pallet": 1}  # train-split snapshot
    loss = report.series_for("loss")
    assert loss is not None and [e for e, _ in loss] == [1, 2, 3]
    assert report.final_metrics["loss"] == pytest.approx(1 / 3)
    assert any(h.name == "epochs" for h in report.hparams)
    assert report.verdict_summary  # the verdict always says something


def test_png_is_one_16_by_9_dashboard(run, tmp_path):
    project, record = run
    path = export_training_report(project, record.run_id, format="png")
    assert path == project.root / "runs" / record.run_id / "exports" / "training_report.png"
    with Image.open(path) as im:
        assert im.size == PNG_SIZE  # 1920x1080
    # an explicit destination is honoured too
    custom = export_training_report(project, record.run_id, format="png",
                                    out_path=tmp_path / "slide.png")
    assert custom.is_file()


def test_pdf_has_the_dashboard_plus_detail_pages(run):
    project, record = run
    path = export_training_report(project, record.run_id, format="pdf")
    data = path.read_bytes()
    assert data.startswith(b"%PDF")
    pages = data.count(b"/Type /Page") - data.count(b"/Type /Pages")
    assert pages >= 4  # dashboard + hyperparameters + evaluation + conclusion


def test_xlsx_sheets(run):
    from openpyxl import load_workbook

    project, record = run
    path = export_training_report(project, record.run_id, format="xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Metrics", "Hyperparameters", "Classes",
                             "Evaluation", "Verdict"]
    summary = {row[0]: row[1] for row in wb["Summary"].iter_rows(min_row=2, values_only=True)}
    assert summary["run_id"] == record.run_id and summary["license"] == "Apache-2.0"
    metrics = list(wb["Metrics"].iter_rows(values_only=True))
    assert metrics[0][0] == "epoch" and "loss" in metrics[0]
    assert [r[0] for r in metrics[1:]] == [1, 2, 3]
    classes = dict(wb["Classes"].iter_rows(min_row=2, values_only=True))
    assert classes == {"forklift": 2, "pallet": 1}


def test_exports_are_listed_newest_first(run):
    project, record = run
    names = {a.name: a for a in list_exports(project, record.run_id)}
    assert {"training_report.png", "training_report.pdf", "training_report.xlsx"} <= set(names)
    assert names["training_report.png"].kind == "report"
    assert names["training_report.xlsx"].format == "xlsx"
    assert all(a.size_bytes > 0 for a in names.values())


def test_unknown_format_and_run_are_refused(run):
    project, record = run
    with pytest.raises(ProjectError, match="Unsupported report format"):
        export_training_report(project, record.run_id, format="docx")
    with pytest.raises(ProjectError, match="No such training run"):
        export_training_report(project, "nope", format="png")
